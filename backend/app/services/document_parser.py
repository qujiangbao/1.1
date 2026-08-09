"""DocumentParser — PDF/DOCX/HTML 文本提取

支持: .pdf (pypdf), .docx (python-docx), .pptx (python-pptx), .html, .txt
"""
from __future__ import annotations

import hashlib
import logging
from pathlib import Path
from dataclasses import dataclass
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    """解析后的文档"""
    raw_text: str
    page_count: int = 1
    file_format: str = ""
    file_hash: str = ""
    file_size: int = 0
    filename: str = ""
    error: Optional[str] = None

    @property
    def is_valid(self) -> bool:
        return bool(self.raw_text) and self.error is None


class DocumentParser:
    """文档解析器"""

    SUPPORTED = {".pdf", ".docx", ".pptx", ".html", ".htm", ".txt", ".md"}

    async def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)
        if not path.exists():
            return ParsedDocument(raw_text="", filename=path.name, error=f"File not found: {file_path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED:
            return ParsedDocument(raw_text="", filename=path.name, error=f"Unsupported format: {ext}")

        file_size = path.stat().st_size
        file_hash = hashlib.sha256(path.read_bytes()).hexdigest()[:16]

        try:
            if ext == ".pdf":
                text, pages = self._parse_pdf(str(path))
            elif ext == ".docx":
                text, pages = self._parse_docx(str(path))
            elif ext == ".pptx":
                text, pages = self._parse_pptx(str(path))
            elif ext in (".html", ".htm"):
                text, pages = self._parse_html(str(path))
            else:
                text = path.read_text(encoding="utf-8")
                pages = 1

            return ParsedDocument(
                raw_text=text.strip(),
                page_count=pages,
                file_format=ext,
                file_hash=file_hash,
                file_size=file_size,
                filename=path.name,
            )
        except Exception as e:
            logger.error(f"Parse failed for {file_path}: {e}")
            return ParsedDocument(raw_text="", filename=path.name, error=str(e))

    def _parse_pdf(self, path: str) -> tuple[str, int]:
        try:
            from pypdf import PdfReader
            doc = PdfReader(path)
            text = "\n\n".join(page.extract_text() or "" for page in doc.pages)
            return text, len(doc.pages)
        except ImportError:
            raise ImportError("需要安装 pypdf 来解析 PDF")

    def _parse_docx(self, path: str) -> tuple[str, int]:
        try:
            from docx import Document
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            return text, 1
        except ImportError:
            raise ImportError("需要安装 python-docx 来解析 DOCX")

    def _parse_pptx(self, path: str) -> tuple[str, int]:
        try:
            from pptx import Presentation
            presentation = Presentation(path)
            slides: list[str] = []
            for slide in presentation.slides:
                parts: list[str] = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            parts.append(" | ".join(cell.text for cell in row.cells))
                slides.append("\n".join(parts))
            return "\n\n".join(slides), len(presentation.slides)
        except ImportError:
            raise ImportError("需要安装 python-pptx 来解析 PPTX")

    def _parse_html(self, path: str) -> tuple[str, int]:
        try:
            from bs4 import BeautifulSoup
            with open(path, encoding="utf-8") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
            return soup.get_text("\n", strip=True), 1
        except ImportError:
            with open(path, encoding="utf-8") as f:
                import re
                text = re.sub(r"<[^>]+>", "", f.read())
            return text, 1
