from zipfile import ZIP_DEFLATED, ZipFile

from app.services.document_parser import DocumentParser


def test_rejects_disguised_office_document(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_text("not a zip archive", encoding="utf-8")

    parsed = DocumentParser()._parse_sync(str(path))

    assert parsed.is_valid is False
    assert "valid Office document" in (parsed.error or "")


def test_rejects_office_archive_with_unsafe_expansion(tmp_path):
    path = tmp_path / "oversized.docx"
    parser = DocumentParser()
    parser.MAX_ARCHIVE_BYTES = 1_024
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (parser.MAX_ARCHIVE_BYTES + 1))

    parsed = parser._parse_sync(str(path))

    assert parsed.is_valid is False
    assert "safe limit" in (parsed.error or "")


def test_rejects_pdf_extension_without_pdf_signature(tmp_path):
    path = tmp_path / "fake.pdf"
    path.write_bytes(b"not-a-pdf")

    parsed = DocumentParser()._parse_sync(str(path))

    assert parsed.is_valid is False
    assert "valid PDF" in (parsed.error or "")


def test_csv_is_converted_to_a_safe_markdown_table(tmp_path):
    path = tmp_path / "enterprises.csv"
    path.write_text(
        "企业名称,统一社会信用代码,注册资本,经营范围\n"
        "广州测试机器人有限公司,91440101TEST000001,1000万元,机器人研发\n",
        encoding="utf-8-sig",
    )

    parsed = DocumentParser()._parse_sync(str(path))

    assert parsed.is_valid is True
    assert "| 企业名称 | 统一社会信用代码 | 注册资本 | 经营范围 |" in parsed.raw_text
    assert "广州测试机器人有限公司" in parsed.raw_text


def test_docx_tables_are_included_in_extracted_text(tmp_path):
    from docx import Document

    path = tmp_path / "enterprise.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "企业名称"
    table.cell(0, 1).text = "统一社会信用代码"
    table.cell(1, 0).text = "广州测试机器人有限公司"
    table.cell(1, 1).text = "91440101TEST000001"
    document.save(path)

    parsed = DocumentParser()._parse_sync(str(path))

    assert parsed.is_valid is True
    assert "广州测试机器人有限公司" in parsed.raw_text
    assert "91440101TEST000001" in parsed.raw_text


def test_xlsx_rows_are_converted_to_markdown(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "enterprises.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "企业台账"
    sheet.append(["企业名称", "统一社会信用代码", "经营范围"])
    sheet.append(["广州测试机器人有限公司", "91440101TEST000001", "机器人研发"])
    workbook.save(path)
    workbook.close()

    parsed = DocumentParser()._parse_sync(str(path))

    assert parsed.is_valid is True
    assert "# 工作表：企业台账" in parsed.raw_text
    assert "广州测试机器人有限公司" in parsed.raw_text
